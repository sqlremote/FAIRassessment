#!/usr/bin/env python3
"""Generate AssessmentRules.xlsx from AssessmentRules.json."""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=14, color="1F3864")
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SEV_FILL = {"blocker": "C00000", "critical": "E26B0A", "high": "FFC000",
            "medium": "BFBFBF", "low": "D9E1F2"}
SEV_FONTCOLOR = {"blocker": "FFFFFF", "critical": "FFFFFF", "high": "000000",
                 "medium": "000000", "low": "000000"}
RESULTS = ["unknown", "pass", "partial", "fail", "na"]

COLS = [
    ("Rule ID", 11), ("Title", 46), ("Severity", 10), ("Type", 13),
    ("Applies to", 22), ("Requirement", 62), ("Why it matters", 62),
    ("Hard limit", 26), ("How to check", 52), ("Remediation", 58),
    ("If not fixed", 50), ("Source", 30), ("Basis", 11),
    ("Result", 11), ("Evidence / link", 26), ("Owner", 16),
    ("Due date", 12), ("Notes", 30),
]
RESULT_COL = 14  # 1-indexed column of "Result"


def fmt_limit(r):
    L = r.get("limit")
    if not L:
        return ""
    v = L["value"]
    if isinstance(v, dict):
        v = "; ".join(f"{k}={vv}" for k, vv in v.items())
    elif isinstance(v, list):
        v = "-".join(str(x) for x in v)
    hard = "hard" if L.get("hard") else "soft"
    return f"{v} {L['unit']} ({hard})"


def fmt_check(c):
    s = c["how"]
    if c.get("operator") and c.get("threshold") is not None:
        s += f"  [{c['operator']} {c['threshold']}{' ' + c['unit'] if c.get('unit') else ''}]"
    elif c.get("expected") is not None and not isinstance(c.get("expected"), (dict, list)):
        s += f"  [expected: {c['expected']}]"
    return f"({c['method']}) {s}"


def style_header(ws, row=1):
    for i, (name, width) in enumerate(COLS, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=4)


def write_rules(ws, rules, sources, start=1, example_row=False):
    style_header(ws, start)
    r = start + 1
    if example_row:
        ex = ["ORG-001", "EXAMPLE ROW - delete before use: your own added rule",
              "high", "config", "copilot-powerbi",
              "What must be true.", "Why it matters for answer quality or cost.",
              "e.g. 200 characters (hard)", "(manual) How to gather the evidence.",
              "What to do about it.", "What goes wrong if you don't.",
              "Your standard", "inferred", "pass", "link to evidence",
              "data.platform@contoso", "2026-10-31", "free text"]
        for i, v in enumerate(ex, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=FONT, size=10, italic=True, color="808080")
            c.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1

    for rule in rules:
        src = "; ".join(sources[s]["title"] for s in rule["sourceIds"] if s in sources)
        vals = [
            rule["id"], rule["title"], rule["severity"], rule["type"],
            ", ".join(rule["appliesTo"]), rule["requirement"], rule["rationale"],
            fmt_limit(rule), fmt_check(rule["check"]), rule["remediation"],
            rule["impactIfViolated"], src, rule["verified"],
            "unknown", "", "", "", rule.get("notes", ""),
        ]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = BODY
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BOX
        ws.cell(row=r, column=1).font = BOLD
        sc = ws.cell(row=r, column=3)
        sc.fill = PatternFill("solid", fgColor=SEV_FILL[rule["severity"]])
        sc.font = Font(name=FONT, size=10, bold=True,
                       color=SEV_FONTCOLOR[rule["severity"]])
        sc.alignment = Alignment(horizontal="center", vertical="center")
        # input cells
        for col in (RESULT_COL, 15, 16, 17, 18):
            ws.cell(row=r, column=col).fill = INPUT_FILL
        ws.row_dimensions[r].height = 62
        r += 1

    last = r - 1
    if last >= start + 1:
        dv = DataValidation(type="list", formula1='"' + ",".join(RESULTS) + '"',
                            allow_blank=False, showDropDown=False)
        dv.error = "Choose one of: " + ", ".join(RESULTS)
        dv.errorTitle = "Invalid result"
        ws.add_data_validation(dv)
        rng = f"{get_column_letter(RESULT_COL)}{start+1}:{get_column_letter(RESULT_COL)}{last}"
        dv.add(rng)
        for val, color, fontcolor in (("pass", "C6EFCE", "006100"),
                                      ("partial", "FFEB9C", "9C5700"),
                                      ("fail", "FFC7CE", "9C0006"),
                                      ("na", "D9D9D9", "595959"),
                                      ("unknown", "FFFFFF", "808080")):
            ws.conditional_formatting.add(rng, CellIsRule(
                operator="equal", formula=[f'"{val}"'],
                fill=PatternFill("solid", fgColor=color),
                font=Font(name=FONT, size=10, bold=True, color=fontcolor)))
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(COLS))}{last}"
    return last


def build(cfg, path):
    wb = Workbook()
    sources = cfg["sources"]
    sev = cfg["severities"]
    by_area = {a["id"]: [r for r in cfg["rules"] if r["area"] == a["id"]]
               for a in cfg["areas"]}
    area_titles = {a["id"]: a["title"] for a in cfg["areas"]}

    # ---------------- Readme
    ws = wb.active
    ws.title = "Readme"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 118
    ws["A1"] = cfg["packName"]; ws["A1"].font = TITLE_FONT
    rows = [
        ("Rule pack version", cfg["packVersion"]),
        ("Spec version", cfg["specVersion"]),
        ("Rules", cfg["ruleCount"]),
        ("Updated", cfg["updated"]),
        ("Purpose", cfg["description"]),
        ("", ""),
        ("HOW TO USE", ""),
        ("1", "Work through one area tab at a time. Each row is one rule."),
        ("2", "Set the yellow Result cell on every row: pass, partial, fail, na, or unknown."),
        ("3", "Record where the evidence lives, who owns the fix, and the target date in the other yellow cells."),
        ("4", "The Scoring tab recalculates live. Read the overall score and maturity level there."),
        ("5", "Any blocker rule left at fail or unknown gates the maturity level at 0 regardless of score."),
        ("", ""),
        ("EDIT ONLY THE YELLOW CELLS", "Every other cell restates the rule and should not be changed here. To change a rule, edit AssessmentRules.json and regenerate this workbook."),
        ("THIS IS A VIEW, NOT THE SOURCE", "AssessmentRules.json is the source of truth. Add your own rules to rules/custom/*.json and rebuild."),
        ("", ""),
        ("RESULT STATES", ""),
    ]
    for k, v in cfg["resultStates"].items():
        credit = "excluded from scoring" if v["credit"] is None else f"credit {v['credit']}"
        rows.append((k, f"{v['description']} ({credit})"))
    rows += [("", ""), ("SEVERITY WEIGHTS", "")]
    for k, v in sev.items():
        rows.append((k, f"weight {v['weight']}{' - GATES the maturity level' if v['gate'] else ''}. {v['description']}"))
    rows += [("", ""), ("BASIS COLUMN", ""),
             ("documented", "Stated in a cited source. Treat the threshold as authoritative but re-verify against current docs."),
             ("inferred", "Engineering judgement built on cited platform behaviour. Calibrate to your own risk appetite."),
             ("community", "Unofficial. Verify before relying on it.")]
    r = 3
    for k, v in rows:
        a, b = ws.cell(row=r, column=1, value=k), ws.cell(row=r, column=2, value=v)
        a.font = BOLD if k and not k.isdigit() else BODY
        b.font = BODY
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if v and len(str(v)) > 90:
            ws.row_dimensions[r].height = 30
        r += 1

    # ---------------- Scoring (live formulas)
    sc = wb.create_sheet("Scoring")
    for col, w in zip("ABCDEFGHIJ", (30, 11, 9, 10, 11, 10, 10, 10, 11, 13)):
        sc.column_dimensions[col].width = w
    sc["A1"] = "Readiness score"; sc["A1"].font = TITLE_FONT
    hdr = ["Area", "Score %", "Weight", "Pass", "Partial", "Fail", "N/A",
           "Unknown", "Earned", "Possible"]
    for i, h in enumerate(hdr, start=1):
        c = sc.cell(row=3, column=i, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    # per-area rows; sheet names are the area titles
    row = 4
    area_rows = {}
    for a in cfg["areas"]:
        rules = by_area[a["id"]]
        if not rules:
            continue
        sheet = a["title"][:31]
        n = len(rules)
        first, last = 3, 2 + n  # header at row 2 (title row 1) on area sheets
        q = f"'{sheet}'!"
        rescol = get_column_letter(RESULT_COL)
        rng = f"{q}${rescol}${first}:${rescol}${last}"
        # weights live in a hidden helper column on each area sheet (col T = 20)
        wcol = f"{q}$T${first}:$T${last}"
        sc.cell(row=row, column=1, value=a["title"]).font = BOLD
        sc.cell(row=row, column=3, value=a["weight"]).font = BODY
        for i, state in enumerate(["pass", "partial", "fail", "na", "unknown"]):
            sc.cell(row=row, column=4 + i,
                    value=f'=COUNTIF({rng},"{state}")').font = BODY
        # earned = 1.0*pass + 0.5*partial ; possible = all except na
        sc.cell(row=row, column=9,
                value=f'=SUMIF({rng},"pass",{wcol})+0.5*SUMIF({rng},"partial",{wcol})').font = BODY
        sc.cell(row=row, column=10,
                value=f'=SUM({wcol})-SUMIF({rng},"na",{wcol})').font = BODY
        sc.cell(row=row, column=2,
                value=f'=IFERROR(ROUND(100*I{row}/J{row},1),"n/a")').font = BOLD
        sc.cell(row=row, column=2).number_format = "0.0"
        area_rows[a["id"]] = row
        row += 1
    last_area = row - 1

    t = row + 1
    sc.cell(row=t, column=1, value="OVERALL (weighted)").font = TITLE_FONT
    sc.cell(row=t, column=2,
            value=f'=IFERROR(ROUND(SUMPRODUCT(B4:B{last_area},C4:C{last_area})'
                  f'/SUM(C4:C{last_area}),1),0)').font = TITLE_FONT
    sc.cell(row=t, column=2).number_format = "0.0"
    for i in range(4, 11):
        sc.cell(row=t, column=i, value=f"=SUM({get_column_letter(i)}4:{get_column_letter(i)}{last_area})").font = BOLD

    # blocker gate
    blockers = [r for r in cfg["rules"] if r["severity"] == "blocker"]
    parts = []
    for b in blockers:
        sheet = area_titles[b["area"]][:31]
        idx = by_area[b["area"]].index(b) + 3
        cell = f"'{sheet}'!{get_column_letter(RESULT_COL)}{idx}"
        parts.append(f'IF(OR({cell}="fail",{cell}="unknown"),1,0)')
    gate = "=" + "+".join(parts) if parts else "=0"
    sc.cell(row=t + 1, column=1, value="Blocker rules not met").font = BOLD
    sc.cell(row=t + 1, column=2, value=gate).font = BOLD

    lv = t + 2
    sc.cell(row=lv, column=1, value="Maturity level").font = TITLE_FONT
    bands = cfg["maturityLevels"]
    nested = f'"{bands[-1]["name"]}"'
    for m in reversed(bands[:-1]):
        nested = f'IF(B{t}<{m["max"]+0.01},"{m["name"]}",{nested})'
    sc.cell(row=lv, column=2,
            value=f'=IF(B{t+1}>0,"{bands[0]["name"]} (GATED)",{nested})').font = TITLE_FONT

    sc.cell(row=lv + 2, column=1, value="Gate rule").font = BOLD
    sc.cell(row=lv + 2, column=2,
            value="Any blocker rule left at fail or unknown forces the maturity level to "
                  f"{bands[0]['name']}, whatever the numeric score. There are "
                  f"{len(blockers)} blocker rules in this pack.").font = NOTE_FONT
    sc.merge_cells(start_row=lv + 2, start_column=2, end_row=lv + 2, end_column=10)
    sc.cell(row=lv + 3, column=1, value="Bands").font = BOLD
    sc.cell(row=lv + 3, column=2, value="  |  ".join(
        f"{m['name']} {m['min']}-{m['max']}" for m in bands)).font = NOTE_FONT
    sc.merge_cells(start_row=lv + 3, start_column=2, end_row=lv + 3, end_column=10)

    for r_ in range(4, last_area + 1):
        sc.conditional_formatting.add(f"B{r_}", CellIsRule(
            operator="lessThan", formula=["60"],
            fill=PatternFill("solid", fgColor="FFC7CE")))
        sc.conditional_formatting.add(f"B{r_}", CellIsRule(
            operator="greaterThanOrEqual", formula=["75"],
            fill=PatternFill("solid", fgColor="C6EFCE")))

    # ---------------- one sheet per area
    for a in cfg["areas"]:
        rules = by_area[a["id"]]
        if not rules:
            continue
        aws = wb.create_sheet(a["title"][:31])
        aws["A1"] = f"{a['title']}  -  {len(rules)} rules, area weight {a['weight']}"
        aws["A1"].font = TITLE_FONT
        aws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
        aws.cell(row=1, column=10,
                 value="Fill the yellow cells only. Result must be one of: "
                       + ", ".join(RESULTS)).font = NOTE_FONT
        write_rules(aws, rules, sources, start=2)
        # hidden weight helper column T
        aws.cell(row=2, column=20, value="_weight").font = NOTE_FONT
        for i, rule in enumerate(rules):
            aws.cell(row=3 + i, column=20,
                     value=sev[rule["severity"]]["weight"] * rule.get("weightMultiplier", 1))
        aws.column_dimensions["T"].hidden = True

    # ---------------- All rules (flat reference)
    allws = wb.create_sheet("All rules")
    allws["A1"] = "All rules - flat reference. Score on the area tabs, not here."
    allws["A1"].font = TITLE_FONT
    write_rules(allws, cfg["rules"], sources, start=2, example_row=True)

    # ---------------- Limits quick reference
    lws = wb.create_sheet("Limits")
    lws["A1"] = "Hard and soft platform limits encoded in this pack"
    lws["A1"].font = TITLE_FONT
    heads = ["Rule ID", "Area", "Limit", "Unit", "Hard?", "What it means",
             "Rule title", "Source"]
    for i, (h, w) in enumerate(zip(heads, (11, 16, 22, 26, 8, 62, 46, 34)), start=1):
        c = lws.cell(row=3, column=i, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(wrap_text=True)
        lws.column_dimensions[get_column_letter(i)].width = w
    r = 4
    for rule in cfg["rules"]:
        L = rule.get("limit")
        if not L:
            continue
        v = L["value"]
        if isinstance(v, dict):
            v = "; ".join(f"{k}={vv}" for k, vv in v.items())
        elif isinstance(v, list):
            v = "-".join(str(x) for x in v)
        vals = [rule["id"], rule["area"], str(v), L["unit"],
                "hard" if L.get("hard") else "soft", L.get("description", ""),
                rule["title"],
                "; ".join(sources[s]["title"] for s in rule["sourceIds"] if s in sources)]
        for i, val in enumerate(vals, start=1):
            c = lws.cell(row=r, column=i, value=val)
            c.font = BODY
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BOX
        lws.row_dimensions[r].height = 30
        r += 1
    lws.auto_filter.ref = f"A3:H{r-1}"
    lws.freeze_panes = "A4"

    # ---------------- Sources
    sws = wb.create_sheet("Sources")
    sws["A1"] = "Sources"; sws["A1"].font = TITLE_FONT
    for i, (h, w) in enumerate(zip(["Source ID", "Title", "URL", "Retrieved", "Note"],
                                   (24, 56, 76, 13, 70)), start=1):
        c = sws.cell(row=3, column=i, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
        sws.column_dimensions[get_column_letter(i)].width = w
    r = 4
    for sid, s in sources.items():
        for i, v in enumerate([sid, s["title"], s.get("url") or "n/a",
                               s.get("retrieved", ""), s.get("note", "")], start=1):
            c = sws.cell(row=r, column=i, value=v)
            c.font = BODY
            c.alignment = Alignment(vertical="top", wrap_text=True)
        if s.get("url"):
            sws.cell(row=r, column=3).hyperlink = s["url"]
            sws.cell(row=r, column=3).font = Font(name=FONT, size=10,
                                                  color="0563C1", underline="single")
        sws.row_dimensions[r].height = 30
        r += 1

    wb.save(path)
    return path


if __name__ == "__main__":
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    cfg = json.load(open(os.path.join(root, "AssessmentRules.json")))
    p = build(cfg, os.path.join(root, "AssessmentRules.xlsx"))
    print("wrote", p)
